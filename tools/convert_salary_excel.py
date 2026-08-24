#!/usr/bin/env python3
"""
Convert salary data from Excel to JSON format.

This script converts an Excel file containing company salary data
into the JSON format expected by salary_lookup.py.

Prerequisites:
    pip install openpyxl

Usage:
    python tools/convert_salary_excel.py <path-to-excel-file>
    python tools/convert_salary_excel.py <path-to-excel-file> --source "My Union Stats 2025"
    python tools/convert_salary_excel.py <path-to-excel-file> --baseline 100 --baseline-desc "Index 100 = median salary"
    python tools/convert_salary_excel.py <path-to-excel-file> --mode absolute --currency INR --unit LPA

The output file (salary_data.json) will be written to the repository root.

Expected Excel format:
    - A header row with column names
    - A "Company" or "Firma" column (required)
    - An optional "City" or "By" column
    - Any number of numeric data columns (salary index, count, etc.)

The script auto-detects the header row and column layout. For Excel files
with paired count/index columns per category, it groups them automatically.

With --mode absolute the data columns are instead classified as money
components by header keyword — Fixed / Base, Variable / Bonus, ESOP / RSU /
Stock, Total / CTC — and grouped into one category per role. Example headers:

    Company | City | AI Engineer 3-5y Count | AI Engineer 3-5y Fixed |
    AI Engineer 3-5y Variable | AI Engineer 3-5y ESOP | AI Engineer 3-5y Total CTC
"""

import json
import sys
import argparse
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


# Column name patterns for auto-detection
COMPANY_PATTERNS = {"firma", "company", "virksomhed", "employer", "arbejdsgiver"}
CITY_PATTERNS = {"by", "city", "kommune", "location", "lokation", "sted"}
COUNT_PATTERNS = {"antal", "count", "number", "n", "employees", "medarbejdere"}
INDEX_PATTERNS = {"indeks", "index", "idx", "salary", "løn", "median", "average", "gennemsnit"}
# "Compound" tokens: pattern words allowed to match as a substring of a larger
# header token, for languages that glue words together (e.g. Danish "lønindeks"
# -> løn + indeks). Languages that write headers as separate words need none.
# Ships populated for this repo's Danish demonstration data; a fork targeting
# another locale edits this constant.
COMPOUND_PATTERNS = {"antal", "indeks", "løn", "gennemsnit", "medarbejdere"}
# Identifier columns (employee id, Danish "personnummer", etc.) are never salary
# data. They are dropped at classification so they are not mistaken for a salary
# category. Matched as whole tokens only, like other pattern sets.
ID_PATTERNS = {"id", "personnummer"}


def parse_numeric_cell(value):
    """Parse numeric Excel values, including localized string cells."""
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        raise ValueError("not numeric")

    text = value.strip().replace("\u00a0", " ").replace(" ", "")
    if not text:
        raise ValueError("not numeric")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        if re.fullmatch(r"[+-]?\d+,\d{3}", text):
            raise ValueError("ambiguous comma separator")
        text = text.replace(",", ".")
    return float(text)


def header_matches(header, patterns):
    """Return True when a header contains a meaningful pattern match.

    Patterns match whole tokens; any pattern also listed in
    ``COMPOUND_PATTERNS`` may additionally match as a substring, to handle
    languages that form compound words.
    """
    h = header.lower().strip()
    tokens = set(re.findall(r"[a-zæøåöäü0-9]+", h))

    for p in patterns:
        if p in tokens:
            return True
        if p in COMPOUND_PATTERNS and p in h:
            return True
    return False


def strip_type_patterns(header, patterns):
    """Remove count/index words from a header to derive a category name."""
    name = header.lower()
    for p in patterns:
        name = re.sub(rf"(?<![a-zæøåöäü0-9]){re.escape(p)}(?![a-zæøåöäü0-9])", "", name)
    return name.strip(" _-")


def detect_column_type(header):
    """Detect whether a column header refers to count or index data."""
    if header_matches(header, COUNT_PATTERNS):
        return "count"
    if header_matches(header, INDEX_PATTERNS):
        return "index"
    return None


# --- Absolute (money) mode --------------------------------------------------
# Header keywords per JSON field. Order matters: the first field whose keywords
# match the header wins, so more specific fields ("total ctc") are checked
# before broader ones.
ABSOLUTE_FIELD_PATTERNS = [
    ("count", COUNT_PATTERNS),
    ("total_ctc_lpa", {"ctc", "total", "package", "totalctc"}),
    ("fixed_lpa", {"fixed", "base", "basic", "fixedpay"}),
    ("variable_lpa", {"variable", "bonus", "incentive", "performance"}),
    ("esop_lpa", {"esop", "esops", "rsu", "rsus", "stock", "equity", "shares"}),
]

ABSOLUTE_FIELDS = ("fixed_lpa", "variable_lpa", "esop_lpa", "total_ctc_lpa")


def detect_absolute_column_type(header):
    """Return the absolute-mode JSON field a column header maps to, or None."""
    for field, patterns in ABSOLUTE_FIELD_PATTERNS:
        if header_matches(header, patterns):
            return field
    return None


def absolute_category_name(header, field):
    """Derive the category key for an absolute-mode column.

    Strips the field keywords out of the header and normalizes what is left:
    "AI Engineer 3-5y Total CTC" -> "ai_engineer_3_5y". A header that is nothing
    but field keywords (e.g. a bare "Fixed" column) falls back to "all".
    """
    patterns = dict(ABSOLUTE_FIELD_PATTERNS)[field]
    name = strip_type_patterns(header, patterns)
    name = re.sub(r"[^a-zæøåöäü0-9]+", "_", name).strip("_")
    return name or "all"


def detect_layout(ws):
    """Locate the header row and the company/city/data columns of a worksheet.

    Returns a dict with keys ``header_row``, ``company_col``, ``city_col`` and
    ``data_cols`` (a list of ``(column_index, header)``), or None when the sheet
    has no usable header. Shared by the index-mode and absolute-mode parsers.
    """
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=False), start=1):
        for cell in row:
            if cell.value and header_matches(str(cell.value), COMPANY_PATTERNS):
                header_row = row_idx
                break
        if header_row:
            break

    if header_row is None:
        print(f"Warning: Could not find header row in sheet '{ws.title}'. Skipping.", file=sys.stderr)
        return None

    headers = []
    for cell in ws[header_row]:
        headers.append(str(cell.value).strip() if cell.value else "")

    company_col = None
    city_col = None
    for i, h in enumerate(headers):
        if header_matches(h, COMPANY_PATTERNS):
            company_col = i
        elif header_matches(h, CITY_PATTERNS):
            city_col = i

    if company_col is None:
        print(f"Warning: Could not find company column in sheet '{ws.title}'.", file=sys.stderr)
        return None

    data_cols = []
    for i, h in enumerate(headers):
        if i == company_col or i == city_col or not h:
            continue
        if header_matches(h, ID_PATTERNS):
            continue
        data_cols.append((i, h))

    return {
        "header_row": header_row,
        "company_col": company_col,
        "city_col": city_col,
        "data_cols": data_cols,
    }


def read_company_city(row, company_col, city_col):
    """Return (company, city) for a data row, or (None, "") when unusable."""
    if company_col >= len(row) or not row[company_col]:
        return None, ""
    company_name = str(row[company_col]).strip()
    if city_col is not None and city_col < len(row) and row[city_col]:
        return company_name, str(row[city_col]).strip()
    return company_name, ""


def parse_sheet_absolute(ws, sheet_label=None):
    """Parse a worksheet into absolute-mode (money) company entries.

    Columns are classified into fixed/variable/esop/total-CTC/count by header
    keywords; the remaining words in the header become the category name, so
    "AI Engineer 3-5y Fixed" and "AI Engineer 3-5y Total CTC" land in the same
    ``ai_engineer_3_5y`` category. Columns that match no absolute field are
    skipped - absolute mode never guesses that an unlabelled number is money.
    """
    layout = detect_layout(ws)
    if layout is None:
        return []

    company_col = layout["company_col"]
    city_col = layout["city_col"]

    # (column_index, category_name, json_field)
    typed_cols = []
    for col_idx, col_header in layout["data_cols"]:
        field = detect_absolute_column_type(col_header)
        if field is None:
            print(
                f"Warning: column '{col_header}' in sheet '{ws.title}' matches no "
                f"absolute-mode field (fixed/variable/esop/total/count). Skipping.",
                file=sys.stderr,
            )
            continue
        cat_name = absolute_category_name(col_header, field)
        typed_cols.append((col_idx, cat_name, field))

    if not typed_cols:
        return []

    companies = []
    for row in ws.iter_rows(min_row=layout["header_row"] + 1, values_only=True):
        company_name, city_name = read_company_city(row, company_col, city_col)
        if company_name is None:
            continue

        entry = {"company": company_name, "city": city_name, "categories": {}}

        for col_idx, cat_name, field in typed_cols:
            if col_idx >= len(row) or row[col_idx] is None:
                continue
            try:
                value = parse_numeric_cell(row[col_idx])
            except (ValueError, TypeError):
                # Non-numeric cell: leave the component absent rather than
                # writing a zero the reader would mistake for a real figure.
                continue
            bucket = entry["categories"].setdefault(cat_name, {})
            bucket[field] = int(value) if field == "count" else value

        entry["categories"] = {k: v for k, v in entry["categories"].items() if v}
        companies.append(entry)

    return companies


def parse_sheet(ws, sheet_label=None):
    """Parse a single worksheet into a list of index-mode company entries."""
    layout = detect_layout(ws)
    if layout is None:
        return []

    header_row = layout["header_row"]
    company_col = layout["company_col"]
    city_col = layout["city_col"]
    data_cols = layout["data_cols"]

    # Group data columns by detected type and derive category names
    count_cols = []
    index_cols = []
    untyped_cols = []

    for col_idx, col_header in data_cols:
        col_type = detect_column_type(col_header)
        if col_type == "count":
            cat_name = strip_type_patterns(col_header, COUNT_PATTERNS)
            count_cols.append((col_idx, col_header, cat_name))
        elif col_type == "index":
            cat_name = strip_type_patterns(col_header, INDEX_PATTERNS)
            index_cols.append((col_idx, col_header, cat_name))
        else:
            untyped_cols.append((col_idx, col_header))

    # Pair count/index columns by matching category name
    categories = []
    used_counts = set()
    used_indexes = set()

    for ci, (c_idx, c_header, c_cat) in enumerate(count_cols):
        for ii, (i_idx, i_header, i_cat) in enumerate(index_cols):
            if ii in used_indexes:
                continue
            if c_cat and i_cat and c_cat == i_cat:
                cat_name = c_cat.replace(" ", "_").replace("-", "_")
                categories.append({
                    "name": cat_name,
                    "count_col": c_idx,
                    "index_col": i_idx,
                })
                used_counts.add(ci)
                used_indexes.add(ii)
                break

    # Remaining unmatched count columns become standalone. They are still count
    # data, so tag them as such — otherwise a lone headcount would be emitted as
    # a salary index and rendered with a meaningless "vs baseline" percentage.
    for ci, (c_idx, c_header, _) in enumerate(count_cols):
        if ci not in used_counts:
            categories.append(
                {"name": c_header.lower().replace(" ", "_"), "value_col": c_idx, "field": "count"}
            )

    # Remaining unmatched index columns become standalone (use original header)
    for ii, (i_idx, i_header, _) in enumerate(index_cols):
        if ii not in used_indexes:
            categories.append({"name": i_header.lower().replace(" ", "_"), "value_col": i_idx})

    # Untyped columns become standalone
    for col_idx, col_header in untyped_cols:
        categories.append({"name": col_header.lower().replace(" ", "_"), "value_col": col_idx})

    # Parse data rows
    companies = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        company_name, city_name = read_company_city(row, company_col, city_col)
        if company_name is None:
            continue

        entry = {
            "company": company_name,
            "city": city_name,
            "categories": {},
        }

        for cat in categories:
            cat_name = cat["name"]
            if "count_col" in cat and "index_col" in cat:
                count_val = None
                index_val = None
                if cat["count_col"] < len(row) and row[cat["count_col"]] is not None:
                    try:
                        count_val = int(parse_numeric_cell(row[cat["count_col"]]))
                    except (ValueError, TypeError):
                        pass
                if cat["index_col"] < len(row) and row[cat["index_col"]] is not None:
                    try:
                        index_val = parse_numeric_cell(row[cat["index_col"]])
                    except (ValueError, TypeError):
                        pass
                # A count/index pair that is entirely empty for this row carries
                # no salary information, so skip it rather than emit nulls.
                if count_val is None and index_val is None:
                    continue
                entry["categories"][cat_name] = {"count": count_val, "index": index_val}
            elif "value_col" in cat:
                if cat["value_col"] < len(row) and row[cat["value_col"]] is not None:
                    val = row[cat["value_col"]]
                    try:
                        val = parse_numeric_cell(val)
                    except (ValueError, TypeError):
                        # Non-numeric standalone value (e.g. a free-text "Notes"
                        # column) is not salary data; skip it for this row.
                        continue
                    field = cat.get("field", "index")
                    entry["categories"][cat_name] = {field: int(val) if field == "count" else val}

        companies.append(entry)

    return companies


def main():
    parser = argparse.ArgumentParser(
        description="Convert salary Excel data to JSON"
    )
    parser.add_argument("excel_file", help="Path to the Excel file with salary data")
    parser.add_argument(
        "--output", default=None,
        help="Output JSON file path (default: salary_data.json in repo root)",
    )
    parser.add_argument(
        "--source", default=None,
        help="Name of the data source (e.g., 'Union Statistics 2025')",
    )
    parser.add_argument(
        "--baseline", type=float, default=100,
        help="Baseline value for index comparison (default: 100)",
    )
    parser.add_argument(
        "--baseline-desc", default=None,
        help="Description of what the baseline means (e.g., 'Index 100 = median salary')",
    )
    parser.add_argument(
        "--mode", choices=["index", "absolute"], default="index",
        help="Output data mode: 'index' (default, index vs a baseline) or "
             "'absolute' (money split into fixed/variable/esop/total CTC)",
    )
    parser.add_argument(
        "--currency", default="INR",
        help="Currency code for --mode absolute (default: INR)",
    )
    parser.add_argument(
        "--unit", default="LPA",
        help="Amount unit for --mode absolute, e.g. LPA or k (default: LPA)",
    )
    args = parser.parse_args()

    excel_path = Path(args.excel_file)
    if not excel_path.exists():
        print(f"Error: File not found: {excel_path}", file=sys.stderr)
        sys.exit(1)

    if openpyxl is None:
        print("Error: openpyxl is required. Install it with: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else Path(__file__).parent.parent / "salary_data.json"

    print(f"Reading: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

    all_companies = []
    for sheet_name in wb.sheetnames:
        print(f"  Parsing sheet: {sheet_name}")
        ws = wb[sheet_name]
        if args.mode == "absolute":
            companies = parse_sheet_absolute(ws, sheet_label=sheet_name)
        else:
            companies = parse_sheet(ws, sheet_label=sheet_name)
        all_companies.extend(companies)

    wb.close()

    if not all_companies:
        print("Error: No data could be parsed from the Excel file.", file=sys.stderr)
        print("Make sure the Excel file has a header row with a 'Company'/'Firma' column.", file=sys.stderr)
        sys.exit(1)

    # Build output
    if args.mode == "absolute":
        metadata = {
            "source": args.source or excel_path.stem,
            "mode": "absolute",
            "currency": args.currency,
            "unit": args.unit,
            "baseline_description": args.baseline_desc
            or f"Absolute compensation in {args.unit} ({args.currency})",
        }
    else:
        metadata = {
            "source": args.source or excel_path.stem,
            "index_baseline": args.baseline,
            "index_label": "Index",
            "baseline_description": args.baseline_desc or f"Index {args.baseline} = baseline",
        }
    output = {"metadata": metadata, "companies": all_companies}

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\nDone! Wrote {len(all_companies)} company entries to {output_path}")


if __name__ == "__main__":
    main()
